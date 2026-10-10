import copy
import math
import os
import re
from decimal import Decimal
from io import BytesIO

from django.conf import settings
from django.core.files.base import ContentFile
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from ..models import AnalysisAttachment, LabAttachmentType, RadiationType
from .jalali import as_date, format_jalali
from .pdf_tools import fix_footer_page_count, tidy_cell_paragraph

TEMPLATE_DIR = os.path.join(settings.BASE_DIR, "laboratory", "templates_docx")

# Dates on every report are printed as Jalali (Solar Hijri). True = Persian digits
# (\u06f1\u06f4\u06f0\u06f5/\u06f0\u06f7/\u06f1\u06f8, the same as the signature rows);
# False = Latin digits (1405/07/18).
JALALI_PERSIAN_DIGITS = True

# One template per unit. All three of a kind share the same row/column layout;
# only the printed unit text differs. NOTE the double underscore in the
# Alpha/Beta kg file name -- that is how the file was uploaded; rename the file
# and this line together if you want it tidied.
GAMMA_TEMPLATES = {
    "KG": "Gamma_analysis_Report_form_kg.docx",
    "L": "Gamma_analysis_Report_form_l.docx",
    "SAMPLE": "Gamma_analysis_Report_form_sample.docx",
}
ALPHA_BETA_TEMPLATES = {
    "KG": "Gross_Alpha___Beta_analysis_Report_form__kg.docx",
    "L": "Gross_Alpha___Beta_analysis_Report_form_l.docx",
    "SAMPLE": "Gross_Alpha___Beta_analysis_Report_form_sample.docx",
}

GAMMA_FIRST_DATA_ROW = 6
GAMMA_DATA_ROWS = 14            # rows 6..19 of the Gamma data table
GAMMA_TABLE_ROWS = 20
ALPHA_BETA_FIRST_DATA_ROW = 6
ALPHA_BETA_ROWS_PER_PAGE = 12   # rows 6..17; a run with more samples gets more pages
ALPHA_BETA_TABLE_ROWS = 23

UNIT_CHOICES = ("KG", "L", "SAMPLE")


def _allow_row_to_grow(cell):
    tr = cell._tc.getparent()
    trPr = tr.find(qn("w:trPr"))
    if trPr is None:
        return
    trHeight = trPr.find(qn("w:trHeight"))
    if trHeight is not None and trHeight.get(qn("w:hRule")) == "exact":
        trHeight.set(qn("w:hRule"), "atLeast")


# =====================================================================
# Persian (RTL) label + value filling -- paragraph bidi, LTR-wrapped numbers.
# =====================================================================

_PPR_SUCCESSORS = {
    "adjustRightInd", "snapToGrid", "spacing", "ind", "contextualSpacing", "mirrorIndents",
    "suppressOverlap", "jc", "textDirection", "textAlignment", "textboxTightWrap",
    "outlineLvl", "divId", "cnfStyle", "rPr", "sectPr", "pPrChange",
}


def _make_paragraph_rtl(paragraph, keep_center=False):
    """Persian form: the paragraph itself must be right-to-left, otherwise Word lays the
    runs out left-to-right and the value lands on the wrong side of the label."""
    pPr = paragraph._p.get_or_add_pPr()
    if pPr.find(qn("w:bidi")) is None:
        bidi = OxmlElement("w:bidi")
        for child in pPr:
            if child.tag.split("}")[1] in _PPR_SUCCESSORS:
                child.addprevious(bidi)
                break
        else:
            pPr.append(bidi)
    jc = pPr.find(qn("w:jc"))
    if jc is not None and not (keep_center and jc.get(qn("w:val")) == "center"):
        pPr.remove(jc)          # in an RTL paragraph the default (start) is already the right edge


# A Latin "island": one or more Latin words / numbers / codes (spaces allowed BETWEEN them, so
# "Test Co" or "12.5 \u00b1 1" stays one unit). A lone " - " is deliberately not part of it.
_LATIN_WORD = r"(?:[A-Za-z0-9][A-Za-z0-9.\-/:\u00b1_]*|\u00b1)"
_LATIN_ISLAND = re.compile(_LATIN_WORD + r"(?:[ \u00a0]+" + _LATIN_WORD + r")*")
_HAS_RTL = re.compile("[\u0590-\u08ff\ufb1d-\ufdff\ufe70-\ufeff]")

LRM = "\u200e"
RLM = "\u200f"


def _add_run(paragraph, text, rtl):
    run = paragraph.add_run(text)
    run.font.name = "B Nazanin"
    run.font.size = Pt(12)
    run.font.rtl = rtl
    rPr = run._r.get_or_add_rPr()
    rPr.find(qn("w:rFonts")).set(qn("w:cs"), "B Nazanin")
    szCs = OxmlElement("w:szCs")
    szCs.set(qn("w:val"), "24")
    rPr.find(qn("w:sz")).addnext(szCs)


def _write_mixed(paragraph, text):
    """Write `text` (Persian and/or Latin) into an RTL paragraph so that it reads in LOGICAL order.

    Why this is more than add_run(text): the bidi algorithm gives neutral characters (space, dash,
    brackets) the direction of the strong text around them. If every Latin chunk is wrapped in LRM
    and the Persian chunks are left bare, the neutrals BETWEEN two Latin chunks fall into the
    left-to-right side and "<date> - <place>" turns into "<place> - <date>". So:
      * every Latin island that contains digits gets LRM on both sides (so 2026-10-05 or AB-20261010-1
        can't be mirrored), and
      * every non-Latin chunk (Persian words, Jalali dates, the " - " between values) gets RLM on both
        sides, which pins those neutrals to the right-to-left side.
    Latin runs are flagged non-RTL, everything else RTL."""
    def rtl_chunk(chunk):
        # keep the direction marks away from the spaces so the line can still break at a space
        core = chunk.strip(" \u00a0")
        if not core:
            _add_run(paragraph, chunk, rtl=True)
            return
        lead = chunk[:len(chunk) - len(chunk.lstrip(" \u00a0"))]
        trail = chunk[len(chunk.rstrip(" \u00a0")):]
        _add_run(paragraph, lead + RLM + core + RLM + trail, rtl=True)

    pos = 0
    for m in _LATIN_ISLAND.finditer(text):
        if m.start() > pos:
            rtl_chunk(text[pos:m.start()])
        island = m.group()
        # LRM only where digits could be mirrored (codes, numbers); plain words keep a clean
        # line-break opportunity, otherwise LibreOffice splits "Test" into "Te" / "st"
        mark = LRM if re.search(r"\d", island) else ""
        _add_run(paragraph, mark + island + mark, rtl=False)
        pos = m.end()
    if pos < len(text):
        rtl_chunk(text[pos:])


def _fill_after_colon(cell, value):
    """Append `value` after the label, keeping the template's own label runs (bold, rtl, fonts) untouched."""
    if value in (None, ""):
        return
    _allow_row_to_grow(cell)
    paragraph = cell.paragraphs[0]
    _make_paragraph_rtl(paragraph)

    text = str(value)
    if not paragraph.text.endswith((" ", "\u00a0")):
        text = " " + text
    _write_mixed(paragraph, text)


def _append_value(cell, value, sep=" "):
    """Plain append for cells whose label is Latin."""
    if value in (None, ""):
        return
    _allow_row_to_grow(cell)
    paragraph = cell.paragraphs[-1] if cell.paragraphs else cell.add_paragraph()
    run = paragraph.add_run(f"{sep}{value}")
    run.font.name = "B Nazanin"
    run.font.size = Pt(12)


def _set_cell_value(cell, value, align=None):
    """Replace a data cell's text. The template's cell paragraphs carry 'space after' and (in the
    No. column) a left indent bigger than the cell can spare, which made two-digit numbers wrap and
    rows grow -- so the paragraph is reset to single spacing, no indent. `align` (e.g.
    WD_ALIGN_PARAGRAPH.CENTER) overrides the template's alignment; None keeps it."""
    if value is None:
        value = ""
    _allow_row_to_grow(cell)
    paragraph = cell.paragraphs[0] if cell.paragraphs else cell.add_paragraph()
    for run in list(paragraph.runs):
        run._r.getparent().remove(run._r)
    text = str(value)
    if _HAS_RTL.search(text):
        # Persian (or mixed) text in a data cell: right-to-left paragraph + the same direction-safe runs
        _make_paragraph_rtl(paragraph, keep_center=True)
        tidy_cell_paragraph(paragraph, align=align)
        _write_mixed(paragraph, text)
        return
    tidy_cell_paragraph(paragraph, align=align)
    run = paragraph.add_run(text)
    run.font.name = "B Nazanin"
    run.font.size = Pt(12)


def _fmt_dt(value):
    """Any date/datetime -> Jalali string (1405/07/18), '' when empty."""
    return format_jalali(value, persian_digits=JALALI_PERSIAN_DIGITS)


def _fmt_num(value):
    """A lab-entered number exactly as entered, minus trailing zeros from the
    DecimalField's fixed 5 places (12.40000 -> 12.4). No rounding, no unit
    conversion: the lab already typed it in the report's unit."""
    if value is None or value == "":
        return ""
    s = format(Decimal(str(value)), "f")
    if "." in s:
        s = s.rstrip("0").rstrip(".")
    return s


def _check_unit(unit):
    unit = (unit or "SAMPLE").upper()
    if unit not in UNIT_CHOICES:
        raise ValueError(f"unit must be one of {UNIT_CHOICES}, got {unit!r}")
    return unit


def _sample_size_cell_text(sample, unit):
    """Value for the Alpha/Beta 'Sample's mass/volume' column, matching what
    each template's own column header says (kg -> mass in kg, l -> volume in
    litres, sample -> volume in ml)."""
    if unit == "KG":
        return _fmt_num(sample.sample_mass_kg)
    if unit == "L":
        return _fmt_num(Decimal(str(sample.sample_volume_ml)) / Decimal(1000)) if sample.sample_volume_ml is not None else ""
    return _fmt_num(sample.sample_volume_ml)


def _replace_attachment(target_kwargs, attachment_type, filename, content, user):
    old = AnalysisAttachment.objects.filter(attachment_type=attachment_type, **target_kwargs).first()
    if old:
        old.file.delete(save=False)
        old.delete()

    attachment = AnalysisAttachment.objects.create(
        attachment_type=attachment_type, uploaded_by=user, **target_kwargs,
    )
    attachment.file.save(filename, ContentFile(content), save=True)
    return attachment


# =====================================================================
# GAMMA -- one Word report per Analysis.
#
# Unit = analysis.effective_result_unit ("KG" / "L" / "SAMPLE"): the unit the
# lab TYPED the results in. Nothing is converted or divided by the sample's
# mass or volume here -- the entered activity, MDA and uncertainty are printed
# exactly as entered, on the template whose headers say that unit.
#
# Layout (verified against the three uploaded templates): 14 per-nuclide rows
# (6..19) with MDA / Uncertainty / Activity / Radionuclide in columns 1-4, and
# two cells that are vertically merged across all of them -- column 0 "Total
# Activity" and column 5 "Code No." -- written once, into the top cell.
# Total Activity is the sum of the Activity values printed in the column, so a
# reader can check it by adding up the page.
# =====================================================================

def generate_gamma_report(analysis, user):

    unit = _check_unit(analysis.effective_result_unit)
    sample = analysis.sample

    activities = list(
        analysis.nuclide_activities.filter(radiation_type=RadiationType.GAMMA).select_related("radionuclide")
    )
    if len(activities) > GAMMA_DATA_ROWS:
        raise ValueError(
            f"This analysis has {len(activities)} gamma nuclides but the Gamma form only has "
            f"{GAMMA_DATA_ROWS} rows. Nothing was dropped silently -- remove or merge some rows."
        )

    doc = Document(os.path.join(TEMPLATE_DIR, GAMMA_TEMPLATES[unit]))

    assert len(doc.tables) == 3, f"Expected 3 tables in the Gamma template, found {len(doc.tables)}."
    assert len(doc.tables[1].rows) == GAMMA_TABLE_ROWS, (
        f"Expected {GAMMA_TABLE_ROWS} rows in the Gamma data table, found {len(doc.tables[1].rows)}."
    )
    assert len(doc.tables[2].rows) == 4, (
        f"Expected 4 rows in the Gamma signature table, found {len(doc.tables[2].rows)}."
    )

    table = doc.tables[1]

    _fill_after_colon(table.cell(1, 0), sample.applicant_name)
    code_bits = " / ".join(filter(None, [sample.sample_type, sample.sample_code_barcode or sample.sample_id]))
    _fill_after_colon(table.cell(1, 3), code_bits)

    location_bits = " - ".join(filter(None, [_fmt_dt(sample.sampling_date), sample.sampling_location]))
    _fill_after_colon(table.cell(2, 0), location_bits)
    _fill_after_colon(table.cell(2, 3), _fmt_dt(analysis.analysis_date))
    _fill_after_colon(table.cell(3, 3), analysis.counting_duration_seconds)

    total = Decimal(0)
    for i, na in enumerate(activities):
        row = GAMMA_FIRST_DATA_ROW + i
        _set_cell_value(table.cell(row, 1), _fmt_num(na.mda_bq))
        _set_cell_value(table.cell(row, 2), _fmt_num(na.uncertainty_bq))
        _set_cell_value(table.cell(row, 3), _fmt_num(na.activity_bq))
        _set_cell_value(table.cell(row, 4), str(na.radionuclide))
        total += Decimal(str(na.activity_bq))

    # merged cells: written once, into the top cell of the merge
    _set_cell_value(table.cell(GAMMA_FIRST_DATA_ROW, 0), _fmt_num(total) if activities else "")
    _set_cell_value(table.cell(GAMMA_FIRST_DATA_ROW, 5), sample.sample_code_barcode or sample.sample_id)

    fix_footer_page_count(doc, total_pages=1)

    buffer = BytesIO()
    doc.save(buffer)
    return _replace_attachment(
        {"analysis": analysis}, LabAttachmentType.REPORT, "gamma_report.docx", buffer.getvalue(), user,
    )


# =====================================================================
# ALPHA/BETA -- one Word report per CountingRun, any number of samples.
#
# Unit = counting_run.result_unit: the lab types the run's alpha/beta results
# and MDAs in that unit (e.g. Bq/Kg, mBq/Kg), and they are printed as entered.
#
# Header block comes from the run's own fields (applicant, sampling
# place/dates, run id, counting duration) -- the same for every sample.
#
# More than 12 samples: the page (header table + data table with its
# signature rows) is repeated, with a page break between copies, inside ONE
# Word document -- so there is still a single file to download, edit and sign.
# Each page carries its own header, MDA row and signature block; the signer
# stamps all of them.
# =====================================================================

def _append_template_pages(doc, extra_pages):
    """Appends `extra_pages` more copies of the template's body content (every
    element except the final section properties).

    Each copy ends with a 'next page' SECTION break (carried by that copy's last, empty
    paragraph) rather than a separate page-break paragraph: when a page is completely full
    a break paragraph is pushed to the next page and leaves a blank page behind it. A
    section break can't do that, and it keeps the same footer on every page."""
    body = doc.element.body
    sect_pr = body.find(qn("w:sectPr"))
    originals = [copy.deepcopy(el) for el in body if el.tag != qn("w:sectPr")]

    def end_section_after(paragraph_el):
        s = copy.deepcopy(sect_pr)
        for old in s.findall(qn("w:type")):
            s.remove(old)
        section_type = OxmlElement("w:type")
        section_type.set(qn("w:val"), "nextPage")
        s.find(qn("w:pgSz")).addprevious(section_type)
        paragraph_el.get_or_add_pPr().append(s)

    last_p = [el for el in body if el.tag == qn("w:p")][-1]
    for _ in range(extra_pages):
        end_section_after(last_p)
        page = [copy.deepcopy(el) for el in originals]
        for el in page:
            sect_pr.addprevious(el)
        last_p = [el for el in page if el.tag == qn("w:p")][-1]


def _date_range_text(first, last):
    first, last = as_date(first), as_date(last)
    if first and last and last != first:
        return f"{_fmt_dt(first)} \u062a\u0627 {_fmt_dt(last)}"      # "<from> \u062a\u0627 <to>"
    return _fmt_dt(first or last)


def generate_alpha_beta_report(counting_run, user):

    unit = _check_unit(counting_run.result_unit)

    analyses = list(counting_run.analyses.select_related("sample").order_by("id"))
    if not analyses:
        raise ValueError("This counting run has no analyses attached \u2014 nothing to report.")

    if counting_run.applicant_name:
        applicant_text = counting_run.applicant_name
    else:
        applicant_names = sorted({a.sample.applicant_name for a in analyses if a.sample.applicant_name})
        applicant_text = "\u060c ".join(applicant_names)

    if counting_run.sampling_location or counting_run.sampling_date_from:
        parts = []
        if counting_run.sampling_date_from:
            parts.append(_date_range_text(counting_run.sampling_date_from, counting_run.sampling_date_to))
        if counting_run.sampling_location:
            parts.append(counting_run.sampling_location)
        sampling_text = " - ".join(parts)
    else:
        dates = sorted({as_date(a.sample.sampling_date) for a in analyses if a.sample.sampling_date})
        places = sorted({a.sample.sampling_location for a in analyses if a.sample.sampling_location})
        parts = []
        if dates:
            parts.append(_date_range_text(dates[0], dates[-1]))
        if places:
            parts.append("\u060c ".join(places))
        sampling_text = " - ".join(parts)

    pages = math.ceil(len(analyses) / ALPHA_BETA_ROWS_PER_PAGE)
    doc = Document(os.path.join(TEMPLATE_DIR, ALPHA_BETA_TEMPLATES[unit]))
    if pages > 1:
        _append_template_pages(doc, pages - 1)

    page_tables = [t for t in doc.tables if len(t.rows) == ALPHA_BETA_TABLE_ROWS]
    assert len(page_tables) == pages, (
        f"Expected {pages} Alpha/Beta data table(s) after building the pages, found {len(page_tables)}. "
        f"Check the template still has exactly one {ALPHA_BETA_TABLE_ROWS}-row table."
    )

    center = WD_ALIGN_PARAGRAPH.CENTER

    for page_index, table in enumerate(page_tables):
        chunk = analyses[page_index * ALPHA_BETA_ROWS_PER_PAGE:(page_index + 1) * ALPHA_BETA_ROWS_PER_PAGE]

        _fill_after_colon(table.cell(1, 0), counting_run.run_id)
        _fill_after_colon(table.cell(1, 4), applicant_text)
        _fill_after_colon(table.cell(2, 0), _fmt_dt(counting_run.run_date))
        _fill_after_colon(table.cell(2, 4), sampling_text)
        _fill_after_colon(table.cell(3, 0), counting_run.counting_duration_seconds)

        for i, analysis in enumerate(chunk):
            row = ALPHA_BETA_FIRST_DATA_ROW + i
            sample = analysis.sample

            _set_cell_value(table.cell(row, 0), str(page_index * ALPHA_BETA_ROWS_PER_PAGE + i + 1), align=center)
            _set_cell_value(table.cell(row, 1), sample.sample_code_barcode or sample.sample_id, align=center)
            _set_cell_value(table.cell(row, 3), _sample_size_cell_text(sample, unit), align=center)

            alpha_text = ""
            if analysis.total_alpha is not None:
                alpha_text = f"{_fmt_num(analysis.total_alpha)} \u00b1 {_fmt_num(analysis.alpha_uncertainty) or '0'}"
            _set_cell_value(table.cell(row, 4), alpha_text, align=center)

            beta_text = ""
            if analysis.total_beta is not None:
                beta_text = f"{_fmt_num(analysis.total_beta)} \u00b1 {_fmt_num(analysis.beta_uncertainty) or '0'}"
            _set_cell_value(table.cell(row, 6), beta_text, align=center)

        # MDA labels in the templates carry their own unit (mBq/Kg, mBq/l, ...);
        # the value is printed as entered, in that unit.
        _append_value(table.cell(18, 0), _fmt_num(counting_run.alpha_mda_mbq))
        _append_value(table.cell(18, 4), _fmt_num(counting_run.beta_mda_mbq))

    fix_footer_page_count(doc, total_pages=pages)

    buffer = BytesIO()
    doc.save(buffer)
    return _replace_attachment(
        {"counting_run": counting_run}, LabAttachmentType.REPORT, "alpha_beta_report.docx", buffer.getvalue(), user,
    )