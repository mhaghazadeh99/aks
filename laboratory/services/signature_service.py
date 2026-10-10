import copy

from django.utils import timezone
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm
from docx.table import _Cell

from operations.services.specification_generator import _format_jalali_date


def _cell_at(table, row_idx, grid_col):
    """Row-aware replacement for table.cell(): python-docx's version assumes
    every row is as wide as the grid, which is false for the Alpha/Beta
    signature rows (7 grid columns' worth of cells in an 8-column grid), so
    cell(row, col) lands one cell off. This walks the row's own cells."""
    tr = table.rows[row_idx]._tr
    pos = tr.grid_before
    for tc in tr.tc_lst:
        span = tc.grid_span
        if pos <= grid_col < pos + span:
            return _Cell(tc, table)
        pos += span
    raise IndexError(f"row {row_idx} has no cell at grid column {grid_col}")


class LabReportSigner:
    """Stamps name / date / signature image into the signature rows of the
    Word report -- the ANALYST / LAB_MANAGER / OPS_MANAGER sequence, same
    floating-picture technique as the license/receive signers.

    Works on the .docx the whole way through the signature chain (people can
    download it, edit it and upload it back between steps); the PDF is only
    made once everyone has signed -- see final_report.py.

    An Alpha/Beta report with more than 12 samples has several pages, each
    with its own signature block (a 23-row table); every one of them is
    stamped. A Gamma report has a single 4-row signature table."""

    GAMMA_SIGNATURE_TABLE_ROWS = 4
    ALPHA_BETA_TABLE_ROWS = 23

    _GAMMA_ROW_BY_ROLE = {"ANALYST": 1, "LAB_MANAGER": 2, "OPS_MANAGER": 3}
    _ALPHA_BETA_ROW_BY_ROLE = {"ANALYST": 20, "LAB_MANAGER": 21, "OPS_MANAGER": 22}

    def __init__(self, document_path, report_kind):
        if report_kind not in ("GAMMA", "ALPHA_BETA"):
            raise ValueError("report_kind must be 'GAMMA' or 'ALPHA_BETA'")
        self.document = Document(document_path)
        self.report_kind = report_kind

    def _add_floating_picture(self, paragraph, image_path, width, top_offset=Mm(0), left_offset=Mm(0)):
        run = paragraph.add_run()
        run.add_picture(image_path, width=width)

        drawing = run._element.find(qn('w:drawing'))
        inline = drawing.find(qn('wp:inline'))

        extent = inline.find(qn('wp:extent'))
        doc_pr = inline.find(qn('wp:docPr'))
        cnv_graphic_frame_pr = inline.find(qn('wp:cNvGraphicFramePr'))
        graphic = inline.find(qn('a:graphic'))

        anchor = OxmlElement('wp:anchor')
        anchor.set('distT', '0'); anchor.set('distB', '0')
        anchor.set('distL', '0'); anchor.set('distR', '0')
        anchor.set('simplePos', '0')
        anchor.set('relativeHeight', '251659264')
        anchor.set('behindDoc', '0')
        anchor.set('locked', '0')
        anchor.set('layoutInCell', '1')
        anchor.set('allowOverlap', '1')

        simple_pos = OxmlElement('wp:simplePos')
        simple_pos.set('x', '0'); simple_pos.set('y', '0')
        anchor.append(simple_pos)

        position_h = OxmlElement('wp:positionH')
        position_h.set('relativeFrom', 'column')
        h_offset = OxmlElement('wp:posOffset')
        h_offset.text = str(int(left_offset))
        position_h.append(h_offset)
        anchor.append(position_h)

        position_v = OxmlElement('wp:positionV')
        position_v.set('relativeFrom', 'paragraph')
        v_offset = OxmlElement('wp:posOffset')
        v_offset.text = str(int(top_offset))
        position_v.append(v_offset)
        anchor.append(position_v)

        anchor.append(copy.deepcopy(extent))

        effect_extent = OxmlElement('wp:effectExtent')
        for side in ('l', 't', 'r', 'b'):
            effect_extent.set(side, '0')
        anchor.append(effect_extent)

        anchor.append(OxmlElement('wp:wrapNone'))
        anchor.append(copy.deepcopy(doc_pr))
        anchor.append(copy.deepcopy(cnv_graphic_frame_pr))
        anchor.append(copy.deepcopy(graphic))

        drawing.replace(inline, anchor)
        return run

    def _stamp(self, signature_cell, date_cell, name_cell, profile):
        name_cell.paragraphs[0].add_run(profile.full_name)
        date_cell.paragraphs[0].add_run(_format_jalali_date(timezone.localdate()))
        signature_path = (
            profile.signature_clean.path
            if profile.signature_clean
            else profile.signature_image.path
        )

        self._add_floating_picture(
            signature_cell.paragraphs[0],
            signature_path,
            width=Mm(28),
            top_offset=Mm(-8),
            left_offset=Mm(0),
        )

    def sign(self, profile, role):
        if not profile.signature_image:
            raise ValueError("User does not have a signature image.")

        if self.report_kind == "GAMMA":
            row = self._GAMMA_ROW_BY_ROLE.get(role)
            if row is None:
                raise ValueError(f"Unknown role: {role}")
            tables = [t for t in self.document.tables if len(t.rows) == self.GAMMA_SIGNATURE_TABLE_ROWS]
            if len(tables) != 1:
                raise ValueError(f"Expected one {self.GAMMA_SIGNATURE_TABLE_ROWS}-row signature table in the Gamma report, found {len(tables)}.")
            table = tables[0]
            # columns: 0 job title | 1 name | 2 date | 3 signature
            self._stamp(
                signature_cell=_cell_at(table, row, 3),
                date_cell=_cell_at(table, row, 2),
                name_cell=_cell_at(table, row, 1),
                profile=profile,
            )
        else:
            row = self._ALPHA_BETA_ROW_BY_ROLE.get(role)
            if row is None:
                raise ValueError(f"Unknown role: {role}")
            tables = [t for t in self.document.tables if len(t.rows) == self.ALPHA_BETA_TABLE_ROWS]
            if not tables:
                raise ValueError("No Alpha/Beta page table found in this report.")
            for table in tables:   # one per page
                # grid: 0-1 signature | 2-3 date | 4 name | 5-6 role (printed)
                self._stamp(
                    signature_cell=_cell_at(table, row, 0),
                    date_cell=_cell_at(table, row, 2),
                    name_cell=_cell_at(table, row, 4),
                    profile=profile,
                )

    def save(self, filename):
        self.document.save(filename)