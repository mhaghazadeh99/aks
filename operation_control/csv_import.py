"""
Reading + validation of import files. No Django imports, so it is testable anywhere.

Accepts, for every part:
  * the ORIGINAL Excel files (.xlsx) - the right sheet is picked automatically, or
  * CSV (UTF-8 / Excel "CSV UTF-8" / legacy Windows-1256), comma, semicolon or tab separated,
with English column names (see TEMPLATES) OR the Persian headers used in your workbooks.
Dates may be Jalali (1405/04/24) or Gregorian (2026-07-15); Persian digits are fine.
"""
import csv
import datetime
import io
import re
from decimal import Decimal

from . import assessment as _assessment
from . import jalali_utils as ju
from .logic import _POSITIONS, to_decimal

SHEETS = {
    "equipment": ["لیست تجهیزات"], "workers": ["پرسنل"], "tld": ["ثبت دز"],
    "online_doses": [], "inspection": ["برنامه سالانه"], "assessments": ["برگه ارزیابی"],
}

# canonical column -> accepted header spellings (Persian ones come from your workbooks)
ALIASES = {
    "equipment": {
        "equipment_id": ["equipment id", "id", "شماره شناسایی تجهیز"],
        "name": ["device name", "device", "نام دستگاه"],
        "model_name": ["model", "مدل"],
        "serial_number": ["serial number", "serial", "شماره سریال"],
        "manufacturer": ["سازنده"],
        "application": ["کاربرد دستگاه", "کاربرد"],
        "status": ["وضعیت دستگاه", "وضعیت"],
        "loan_to": ["loaned to"],
        "center": ["مرکز نگهداری"],
        "current_location": ["current location", "موقعیت دستگاه در حال حاضر"],
        "calibration_period_months": ["period months", "period", "دوره کالیبراسیون ماه", "دوره کالیبراسیون"],
        "current_label": ["label", "برچسب"],
        "notes": ["یادداشت اقدام", "یادداشت"],
        "last_calibration_date": ["last calibration", "calibration date", "تاریخ آخرین کالیبراسیون شمسی", "تاریخ آخرین کالیبراسیون"],
        "certificate_number": ["certificate number", "certificate", "شماره گواهی کالیبراسیون"],
        "laboratory": ["lab", "آزمایشگاه مرجع کالیبراسیون"],
        "calibration_type": ["calibration type"],
    },
    "workers": {
        "personnel_code": ["code", "کد پرسنلی شماره فرد", "کد پرسنلی"],
        "full_name": ["name", "نام و نام خانوادگی"],
        "unit": ["واحد بخش", "واحد"],
        "position": ["سمت"],
        "status": ["وضعیت"],
        "exit_date": ["تاریخ خروج شمسی", "تاریخ خروج"],
        "notes": ["توضیحات"],
    },
    "tld": {
        "year": ["jalali year", "سال شمسی دوره", "سال"],
        "season": ["فصل دوره", "فصل"],
        "report_number": ["report number", "شماره گزارش دریافتی", "شماره گزارش"],
        "received_date": ["report date", "date received", "تاریخ ثبت گزارش شمسی", "تاریخ ثبت گزارش"],
        "tld_code": ["tld", "کد tld"],
        "personnel_code": ["code", "worker", "کد پرسنلی شماره فرد", "کد پرسنلی"],
        "full_name": ["name", "نام و نام خانوادگی"],
        "env_point": ["environmental point", "نقطه محیطی"],
        "position": ["wearing position", "محل پوشش"],
        "gamma_msv": ["gamma", "دز گاما"],
        "neutron_msv": ["neutron", "دز نوترون"],
        "dose": ["مقدار دز"],
        "unit": ["واحد"],
        "remarks": ["notes", "توضیحات"],
    },
    "online_doses": {
        "date": ["operation date", "performed date", "تاریخ عملیات"],
        "operation": ["operation type", "عملیات", "نوع عملیات"],
        "location": ["محل"],
        "operation_duration_min": ["duration minutes", "operation duration", "مدت عملیات دقیقه"],
        "description": ["شرح عملیات"],
        "session": ["session id", "شناسه عملیات"],
        "personnel_code": ["code", "worker", "کد پرسنلی"],
        "full_name": ["name", "نام و نام خانوادگی"],
        "device": ["device id", "equipment id", "دستگاه", "شماره شناسایی تجهیز"],
        "dose_type": ["نوع دز"],
        "dose_value": ["dose", "مقدار دز"],
        "unit": ["واحد"],
        "exposure_minutes": ["received over min", "مدت دریافت دقیقه"],
    },
    "assessments": {
        "equipment_id": ["equipment id", "id", "device", "نام کد تجهیز حفاظت پرتوی"],
        "assessment_date": ["date", "تاریخ ارزیابی"],
        "assessed_by": ["assessor", "ارزیاب"],
        "has_maker_recommendation": ["پیشنهاد سازنده موجود", "maker recommendation available", "has maker recommendation"],
        "device_type": ["نوع دستگاه"],
        "maker_recommendation": ["پیشنهاد کارخانه سازنده", "manufacturer recommendation", "maker recommendation"],
        "accuracy": ["دقت اندازه گیری msa", "دقت اندازه گیری", "msa"],
        "environment": ["تاثیر شرایط محیطی"],
        "importance": ["حوزه استفاده و اهمیت"],
        "wear": ["فرسودگی عوامل مخرب"],
        "cost": ["هزینه کالیبراسیون"],
        "notes": ["ملاحظات"],
    },
    "inspection": {
        "title": ["inspection", "عنوان بازرسی"],
        "form_title": ["form", "چک لیست مربوطه فرم مربوطه"],
        "form_code": ["کد فرم"],
        "area": ["نواحی", "ناحیه"],
        "period": ["period weeks", "دوره"],
        "responsible": ["مسئول انجام"],
        "is_active": ["active"],
        "done_weeks": ["done weeks"],
        "plan_year": ["year"],
    },
}

REQUIRED = {
    "equipment": ["equipment_id", "name"],
    "workers": ["personnel_code", "full_name"],
    "tld": ["year", "season", "tld_code"],
    "online_doses": ["date", "operation", "device", "dose_type", "dose_value"],
    "inspection": ["title", "area", "period"],
    "assessments": ["equipment_id", "device_type", "accuracy", "environment", "importance", "wear", "cost"],
}

TEMPLATES = {
    "equipment": (
        ["equipment_id", "name", "model_name", "serial_number", "manufacturer", "application", "status", "center",
         "current_location", "calibration_period_months", "current_label", "notes", "last_calibration_date",
         "certificate_number", "laboratory", "calibration_type"],
        [["PH-0001E1", "GR-135", "GR-135ND-H", "3399CN", "SAIC Expioranium", "Survey meter", "Active", "Tehran",
          "Health physics room", "12", "Calibrated", "", "1402/09/15", "02091564", "SSDL", "External"]]),
    "workers": (
        ["personnel_code", "full_name", "unit", "position", "status", "exit_date", "notes"],
        [["2050625911", "Ali Asghari", "Operations", "Technician", "Active", "", ""]]),
    "tld": (
        ["year", "season", "report_number", "received_date", "tld_code", "personnel_code", "full_name", "env_point",
         "position", "gamma_msv", "neutron_msv", "unit", "remarks"],
        [["1404", "Spring", "TL6289-0412", "1405/04/21", "TL-001", "2050625911", "Ali Asghari", "", "body", "0.35", "0.02", "mSv", ""],
         ["1404", "Spring", "TL6289-0412", "1405/04/21", "ENV-01", "", "", "Main gate", "environmental", "0.12", "", "mSv", ""]]),
    "online_doses": (
        ["date", "operation", "location", "operation_duration_min", "description", "personnel_code", "device",
         "dose_type", "dose_value", "unit", "exposure_minutes"],
        [["1405/07/20", "Waste conditioning", "Tehran hall", "90", "", "2050625911", "PH-0007E21", "gamma", "35", "uSv", "60"],
         ["1405/07/20", "Waste conditioning", "Tehran hall", "90", "", "1590195329", "PH-0007E22", "neutron", "4", "uSv", "60"]]),
    "assessments": (
        ["equipment_id", "assessment_date", "assessed_by", "has_maker_recommendation", "device_type",
         "maker_recommendation", "accuracy", "environment", "importance", "wear", "cost", "notes"],
        [["PH-0001E2", "1405/07/20", "2050625911", "yes", "portable", "one_year", "high", "medium", "high", "high", "very_high", ""],
         ["PH-0010", "1405/07/20", "", "no", "fixed", "", "medium", "good", "medium", "low", "low", ""]]),
    "inspection": (
        ["title", "form_title", "form_code", "area", "period", "responsible", "is_active", "done_weeks"],
        [["Environmental dose rate monitoring", "Dose monitoring form - Tehran", "S4-FRM-16/01", "Tehran treatment hall",
          "Every two weeks", "", "yes", "1,3,5"]]),
}

_DIGITS = str.maketrans("۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩", "01234567890123456789")


class ImportFormatError(Exception):
    """The file cannot be read at all (bad format / unrecognised columns)."""


def normalize_digits(s):
    return str(s).translate(_DIGITS)


def norm_header(h):
    s = str(h if h is not None else "")
    s = s.replace("\u200c", " ").replace("ي", "ی").replace("ك", "ک").replace("ـ", "")
    return re.sub(r"[\s_\-/()\[\].:؟?،؛]+", " ", s).strip().lower()


def _reverse(kind):
    rev = {}
    for canon, aliases in ALIASES[kind].items():
        for a in [canon, *aliases]:
            rev.setdefault(norm_header(a), canon)
    return rev


def cell_text(v):
    if v is None:
        return ""
    if isinstance(v, (datetime.date, datetime.datetime)):
        return v
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return " ".join(str(v).replace("ي", "ی").replace("ك", "ک").split())


# ------------------------------------------------------------------ raw readers
def _decode(data: bytes) -> str:
    if data[:2] in (b"\xff\xfe", b"\xfe\xff"):
        return data.decode("utf-16")
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("cp1256", errors="replace")      # Excel "CSV" with Persian text on Windows


def _csv_rows(data):
    text = _decode(data)
    first = next((ln for ln in text.splitlines() if ln.strip()), "")
    delim = max([",", ";", "\t"], key=first.count)
    reader = csv.reader(io.StringIO(text), delimiter=delim)
    return [(n, row) for n, row in enumerate(reader, 1) if any(str(c).strip() for c in row)]


def _xlsx_rows(data, kind, extras):
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheet = next((wb[n] for n in SHEETS[kind] if n in wb.sheetnames), wb.worksheets[0])
    rows, empty = [], 0
    for n, r in enumerate(sheet.iter_rows(values_only=True), 1):
        if r is None or all(v in (None, "") for v in r):
            empty += 1
            if empty > 60:                      # the used range of these sheets runs to ~100k rows
                break
            continue
        empty = 0
        rows.append((n, list(r)))
    if kind == "equipment" and "Sheet1" in wb.sheetnames:       # current labels live on Sheet1 in your workbook
        labels = {}
        for r in wb["Sheet1"].iter_rows(min_row=2, values_only=True):
            if r and len(r) > 7 and r[1] and r[7]:
                labels[str(cell_text(r[1]))] = cell_text(r[7])
        extras["labels"] = labels
    return rows


def load_table(filename, data, kind):
    name = (filename or "").lower()
    extras = {}
    if name.endswith((".xlsx", ".xlsm")):
        raw = _xlsx_rows(data, kind, extras)
    elif name.endswith(".xls"):
        raise ImportFormatError("Old .xls files are not supported - save the file as .xlsx or CSV UTF-8 first.")
    else:
        raw = _csv_rows(data)
    rev = _reverse(kind)

    best_i, best, best_n = None, None, 0
    for i, (_n, values) in enumerate(raw[:12]):
        canon = [rev.get(norm_header(c)) if c not in (None, "") else None for c in values]
        n = len({c for c in canon if c})
        if n > best_n:
            best_i, best, best_n = i, canon, n
    if best_n < 2:
        raise ImportFormatError("Could not find the header row. Expected columns like: "
                                + ", ".join(TEMPLATES[kind][0]))

    rows, present = [], {c for c in best if c}
    for n, values in raw[best_i + 1:]:
        row, seen = {"_row": n}, set()
        for canon, v in zip(best, values):
            if canon and canon not in seen:
                row[canon] = cell_text(v)
                seen.add(canon)
        if any(v not in ("", None) for k, v in row.items() if k != "_row"):
            rows.append(row)
    if kind == "equipment" and extras.get("labels") and "current_label" not in present:
        for row in rows:
            if extras["labels"].get(row.get("equipment_id", "")):
                row["current_label"] = extras["labels"][row["equipment_id"]]
        present.add("current_label")
    return {"rows": rows, "present": present, "extras": extras}


def missing_columns(kind, present):
    missing = [c for c in REQUIRED[kind] if c not in present]
    if kind == "tld":
        if not present & {"personnel_code", "full_name", "env_point"}:
            missing.append("personnel_code / full_name / env_point")
        if not present & {"gamma_msv", "neutron_msv", "dose"}:
            missing.append("gamma_msv / neutron_msv (or dose)")
    if kind == "online_doses" and not present & {"personnel_code", "full_name"}:
        missing.append("personnel_code / full_name")
    return missing


# ------------------------------------------------------------------ value parsers
def parse_date(v):
    """Jalali (year < 1700) or Gregorian; real Excel dates pass through. '' -> None; bad -> ValueError."""
    if v in (None, ""):
        return None
    if isinstance(v, datetime.datetime):
        return v.date()
    if isinstance(v, datetime.date):
        return v
    m = re.match(r"^(\d{4})[/\-.](\d{1,2})[/\-.](\d{1,2})$", normalize_digits(v).strip())
    if not m:
        raise ValueError(v)
    y, mo, d = (int(g) for g in m.groups())
    if y < 1700:
        ymd = ju.parse_jalali_str(f"{y}/{mo}/{d}")
        if not ymd:
            raise ValueError(v)
        return ju.from_jalali(*ymd)
    return datetime.date(y, mo, d)


def _int(v):
    try:
        return int(normalize_digits(v).strip())
    except (ValueError, TypeError):
        raise ValueError(v)


def _lookup(table, value):
    return table.get(norm_header(value))


def _t(row, key):
    v = row.get(key, "")
    return v if isinstance(v, str) else str(v)


STATUS_EQ = {"فعال": "ACTIVE", "active": "ACTIVE", "غیرفعال": "INACTIVE", "inactive": "INACTIVE",
             "ارسال به کالیبراسیون": "SENT", "sent": "SENT", "sent for calibration": "SENT",
             "on loan": "LOAN", "loan": "LOAN"}
CENTER = {"تهران": "TEHRAN", "tehran": "TEHRAN", "کرج": "KARAJ", "karaj": "KARAJ", "انارک": "ANARAK", "anarak": "ANARAK"}
LABEL = {"calibrated": "CALIBRATED", "rejected": "REJECTED", "out of service": "OUT_OF_SERVICE", "cbu": "CBU",
         "تصدیق": "VERIFICATION", "verification": "VERIFICATION"}
CAL_TYPE = {"external": "EXTERNAL", "خارجی": "EXTERNAL", "int ref": "INT_REF", "internal with ref": "INT_REF",
            "internal": "INT_REF", "داخلی": "INT_REF", "int noref": "INT_NOREF", "internal no ref": "INT_NOREF"}
SEASON = {"بهار": 1, "تابستان": 2, "پاییز": 3, "زمستان": 4, "spring": 1, "summer": 2, "autumn": 3, "fall": 3,
          "winter": 4, "1": 1, "2": 2, "3": 3, "4": 4}
DOSE_TYPE = {"gamma": "GAMMA", "گاما": "GAMMA", "γ": "GAMMA", "neutron": "NEUTRON", "نوترون": "NEUTRON", "n": "NEUTRON"}
PERIOD = {"هفتگی": 1, "weekly": 1, "دوهفته یکبار": 2, "دو هفته یکبار": 2, "biweekly": 2, "every two weeks": 2,
          "ماهانه": 4, "monthly": 4, "دو ماهه": 8, "bimonthly": 8, "every two months": 8,
          "سه ماهه": 12, "quarterly": 12, "every three months": 12,
          "شش ماهه": 24, "semiannual": 24, "every six months": 24}
_YES = {"1", "yes", "true", "y", "بله", "فعال", "active"}
_NO = {"0", "no", "false", "n", "خیر", "غیرفعال", "inactive"}


def _unit_factor_msv(text):
    t = norm_header(text).replace("µ", "u").replace("μ", "u")
    if t == "":
        return 1
    if t in ("msv", "میلی سیورت", "میلیسیورت"):
        return 1
    if t in ("usv", "میکرو سیورت", "میکروسیورت"):
        return 0.001
    if t in ("sv", "سیورت"):
        return 1000
    raise ValueError(text)


def _mk(errors, warnings):
    return (lambda r, m: errors.append((r, m))), (lambda r, m: warnings.append((r, m)))


# ------------------------------------------------------------------ parsers (one per part)
def parse_equipment(table):
    items, errors, warnings = [], [], []
    err, warn = _mk(errors, warnings)
    present, seen = table["present"], set()
    for row in table["rows"]:
        rn, eid, name = row["_row"], _t(row, "equipment_id"), _t(row, "name")
        if not eid:
            err(rn, f"missing equipment ID ('{name}')")
            continue
        if eid in seen:
            warn(rn, f"{eid} appears more than once - the later row updates the earlier one")
        seen.add(eid)
        f, ok = {}, True
        for key in ("name", "model_name", "serial_number", "manufacturer", "current_location", "notes", "loan_to"):
            if key in present:
                f[key] = _t(row, key)
        if "name" in present and not f["name"]:
            warn(rn, f"{eid}: empty device name")
        if "status" in present:
            raw = _t(row, "status")
            if "امانت" in raw:
                f["status"], f["loan_to"] = "LOAN", raw.replace("امانت", "", 1).replace("به", "", 1).strip()
            else:
                f["status"] = _lookup(STATUS_EQ, raw) if raw else "ACTIVE"
                if not f["status"]:
                    warn(rn, f"{eid}: unknown status '{raw}' - set to Active")
                    f["status"] = "ACTIVE"
        if "امانت" in f.get("current_location", ""):
            f["current_location"] = ""
        if "center" in present:
            raw = _t(row, "center")
            f["center"] = _lookup(CENTER, raw) or ""
            if raw and not f["center"]:
                warn(rn, f"{eid}: center '{raw}' is not Tehran/Karaj/Anarak - left blank")
        if "current_label" in present:
            raw = _t(row, "current_label")
            f["current_label"] = _lookup(LABEL, raw) or ""
            if raw and not f["current_label"]:
                warn(rn, f"{eid}: unknown label '{raw}' - left blank")
        if "calibration_period_months" in present:
            raw = _t(row, "calibration_period_months")
            months = None
            try:
                months = _int(raw) if raw not in ("", "-") else 0
            except ValueError:
                # your sheet has text here (e.g. "after gas replacement"): keep it as a note, default to 12
                note = f"{raw}"
                f["notes"] = (f.get("notes", "") + " | " if f.get("notes") else "") + f"دوره کالیبراسیون / calibration period: {note}"
                warn(rn, f"{eid}: calibration period '{raw}' is not a number - set to 12 months and copied to the notes")
                months = 12
            if months < 0:
                err(rn, f"{eid}: invalid calibration period '{raw}'")
                ok = False
            else:
                if months == 0:
                    warn(rn, f"{eid}: calibration period is {raw or 'empty'} - set to 12 months, please check")
                    months = 12
                f["calibration_period_months"] = months
        cal = None
        if "last_calibration_date" in present and row.get("last_calibration_date") not in ("", None):
            try:
                cal = {"date": parse_date(row["last_calibration_date"])}
                for key in ("certificate_number", "laboratory"):
                    cal[key] = _t(row, key) if key in present else ""
                if "calibration_type" in present and _t(row, "calibration_type"):
                    cal["calibration_type"] = _lookup(CAL_TYPE, _t(row, "calibration_type")) or ""
            except ValueError:
                err(rn, f"{eid}: invalid calibration date '{row['last_calibration_date']}'")
                ok = False
        if ok:
            items.append({"row": rn, "equipment_id": eid, "fields": f, "calibration": cal,
                          "application": _t(row, "application") if "application" in present else ""})
    return {"items": items, "errors": errors, "warnings": warnings}


def parse_workers(table):
    items, errors, warnings = [], [], []
    err, warn = _mk(errors, warnings)
    present = table["present"]
    for row in table["rows"]:
        rn, code, name = row["_row"], _t(row, "personnel_code"), _t(row, "full_name")
        if not code and not name:
            continue                                    # stray / pre-filled row (e.g. only a default date)
        if not code or not name:
            err(rn, "personnel code and full name are both required")
            continue
        f = {"full_name": name}
        for key in ("unit", "position", "notes"):
            if key in present:
                f[key] = _t(row, key)
        if "status" in present:
            raw = norm_header(_t(row, "status"))
            if not raw or "فعال" in raw and "غیر" not in raw or raw == "active":
                f["status"] = "ACTIVE"
            elif any(w in raw for w in ("غیر", "خارج", "ترک", "بازنشست", "منتقل", "left", "inactive")):
                f["status"] = "LEFT"
            else:
                warn(rn, f"{code}: unknown status '{raw}' - set to Active")
                f["status"] = "ACTIVE"
        if "exit_date" in present:
            try:
                f["exit_date"] = parse_date(row.get("exit_date"))
            except ValueError:
                err(rn, f"{code}: invalid exit date '{row.get('exit_date')}'")
                continue
        items.append({"row": rn, "personnel_code": code, "fields": f})
    return {"items": items, "errors": errors, "warnings": warnings}


def parse_tld(table):
    items, errors, warnings = [], [], []
    err, warn = _mk(errors, warnings)
    present = table["present"]
    legacy = "dose" in present and not present & {"gamma_msv", "neutron_msv"}
    if legacy:
        warnings.append((0, "A single 'dose' column was found - it is imported as GAMMA dose "
                            "(add gamma_msv / neutron_msv columns to separate them)."))
    for row in table["rows"]:
        rn = row["_row"]
        code, name, env, tld = (_t(row, k) for k in ("personnel_code", "full_name", "env_point", "tld_code"))
        if not (code or name or env or tld or _t(row, "gamma_msv") or _t(row, "neutron_msv") or _t(row, "dose")):
            continue                                   # pre-filled template row with no data
        try:
            year = _int(row.get("year", ""))
            season = _lookup(SEASON, _t(row, "season"))
            if not season:
                raise ValueError(_t(row, "season"))
            received = parse_date(row.get("received_date")) if "received_date" in present else None
        except ValueError as e:
            err(rn, f"invalid year / season / date: '{e}'")
            continue
        if not tld:
            err(rn, "TLD code is missing")
            continue
        pos_raw = _t(row, "position").lower()
        position = _POSITIONS.get(pos_raw) or ("ENV" if env else "BODY")
        if pos_raw and pos_raw not in _POSITIONS:
            err(rn, f"unknown position '{pos_raw}' (body / hand / environmental)")
            continue
        if position == "ENV" and not env:
            err(rn, "environmental TLD needs an env_point")
            continue
        if position != "ENV" and not (code or name):
            err(rn, "body/hand TLD needs a personnel code or name")
            continue
        try:
            factor = _unit_factor_msv(_t(row, "unit")) if "unit" in present else 1
            gamma = to_decimal(_t(row, "gamma_msv")) if "gamma_msv" in present else None
            neutron = to_decimal(_t(row, "neutron_msv")) if "neutron_msv" in present else None
            if legacy or (gamma is None and "dose" in present):
                gamma = to_decimal(_t(row, "dose")) if gamma is None else gamma
        except ValueError as e:
            err(rn, f"invalid dose or unit: '{e}'")
            continue
        if gamma is None and neutron is None:
            err(rn, "no dose value")
            continue
        f = Decimal(str(factor))
        items.append({"row": rn, "year": year, "season": season, "report_number": _t(row, "report_number"),
                      "received_date": received, "tld_code": tld, "position": position, "code": code, "name": name,
                      "env": env, "gamma": gamma * f if gamma is not None else None,
                      "neutron": neutron * f if neutron is not None else None, "remarks": _t(row, "remarks")})
    return {"items": items, "errors": errors, "warnings": warnings}


def parse_online_doses(table):
    items, errors, warnings = [], [], []
    err, warn = _mk(errors, warnings)
    present = table["present"]
    for row in table["rows"]:
        rn = row["_row"]
        try:
            date = parse_date(row.get("date"))
            if date is None:
                raise ValueError("date")
            dose = to_decimal(_t(row, "dose_value"))
            if dose is None:
                raise ValueError("dose")
            raw_unit = _t(row, "unit")
            factor = _unit_factor_msv(raw_unit) if raw_unit else 0.001      # online devices: blank = uSv
            unit = "MSV" if factor == 1 else "USV"
            if factor not in (1, 0.001):
                dose, unit = dose * 1000, "MSV"          # given in Sv
            duration = _int(_t(row, "operation_duration_min")) if _t(row, "operation_duration_min") else None
            exposure = _int(_t(row, "exposure_minutes")) if _t(row, "exposure_minutes") else None
        except ValueError as e:
            err(rn, f"invalid date, dose, unit or minutes: '{e}'")
            continue
        dose_type = _lookup(DOSE_TYPE, _t(row, "dose_type"))
        operation, device = _t(row, "operation"), _t(row, "device")
        code, name = _t(row, "personnel_code"), _t(row, "full_name")
        if not dose_type:
            err(rn, f"unknown dose type '{_t(row, 'dose_type')}' (gamma / neutron)")
        elif not (operation and device and (code or name)):
            err(rn, "operation, device and worker are required")
        else:
            items.append({"row": rn, "date": date, "operation": operation, "location": _t(row, "location"),
                          "duration": duration, "description": _t(row, "description"), "session": _t(row, "session"),
                          "code": code, "name": name, "device": device, "dose_type": dose_type,
                          "dose": dose, "unit": unit, "exposure": exposure})
    return {"items": items, "errors": errors, "warnings": warnings}


def _weeks(text):
    out = []
    for part in re.split(r"[,\s،;]+", normalize_digits(text).strip()):
        if not part:
            continue
        if "-" in part:
            a, b = part.split("-", 1)
            out.extend(range(int(a), int(b) + 1))
        else:
            out.append(int(part))
    return out


def parse_inspection(table):
    items, errors, warnings = [], [], []
    err, warn = _mk(errors, warnings)
    present, title = table["present"], ""
    for row in table["rows"]:
        rn = row["_row"]
        title = _t(row, "title") or title                 # blank title = same as the row above (as in your Excel)
        area, raw_period = _t(row, "area"), _t(row, "period")
        if not title or not area or not raw_period:
            err(rn, "title, area and period are required")
            continue
        period = _lookup(PERIOD, raw_period)
        if period is None:
            try:
                period = _int(raw_period)
            except ValueError:
                period = None
        if period not in (1, 2, 4, 8, 12, 24):
            err(rn, f"unknown period '{raw_period}' (weekly / every two weeks / monthly / ...)")
            continue
        f = {"period_weeks": period}
        for key in ("form_title", "form_code"):
            if key in present:
                f[key] = _t(row, key)
        if "is_active" in present and _t(row, "is_active"):
            v = norm_header(_t(row, "is_active"))
            f["is_active"] = True if v in _YES else False if v in _NO else True
        try:
            done = _weeks(_t(row, "done_weeks")) if "done_weeks" in present else []
            year = _int(_t(row, "plan_year")) if "plan_year" in present and _t(row, "plan_year") else None
        except ValueError as e:
            err(rn, f"invalid done weeks / plan year: '{e}'")
            continue
        items.append({"row": rn, "title": title, "area": area, "fields": f, "responsible": _t(row, "responsible"),
                      "done_weeks": done, "plan_year": year})
    return {"items": items, "errors": errors, "warnings": warnings}


def _option(factor, text):
    n = norm_header(text)
    for code, row in _assessment.TABLE[factor].items():
        if n in (norm_header(code), norm_header(row["fa"]), norm_header(row["en"])):
            return code
    return None


def parse_assessments(table):
    items, errors, warnings = [], [], []
    err, warn = _mk(errors, warnings)
    present = table["present"]
    for row in table["rows"]:
        rn, ref = row["_row"], _t(row, "equipment_id")
        if not ref:
            err(rn, "device name / code is missing")
            continue
        if "نمونه تکمیل" in _t(row, "notes"):
            warn(rn, f"{ref}: skipped - the sheet marks this row as a sample ('نمونه تکمیل‌شده')")
            continue
        maker_raw = _t(row, "maker_recommendation")
        raw_has = norm_header(_t(row, "has_maker_recommendation"))
        has_maker = True if raw_has in _YES or raw_has == "بله" else False if raw_has in _NO or raw_has == "خیر" \
            else bool(maker_raw)
        problems, answers = [], {}
        for factor, key in (("device_type", "device_type"), ("accuracy", "accuracy"), ("environment", "environment"),
                            ("importance", "importance"), ("wear", "wear"), ("cost", "cost")):
            code = _option(factor, _t(row, key))
            if code is None:
                problems.append(f"{key} '{_t(row, key)}'")
            answers[key] = code
        maker = None
        if has_maker:
            maker = _option("maker", maker_raw)
            if maker is None:
                problems.append(f"maker_recommendation '{maker_raw}'")
        if problems:
            err(rn, f"{ref}: missing or unknown answer(s): " + ", ".join(problems))
            continue
        try:
            date = parse_date(row.get("assessment_date")) if "assessment_date" in present else None
        except ValueError:
            err(rn, f"{ref}: invalid date '{row.get('assessment_date')}'")
            continue
        items.append({"row": rn, "ref": ref, "date": date, "assessor": _t(row, "assessed_by"),
                      "has_maker": has_maker, "maker": maker, "notes": _t(row, "notes"), **answers})
    return {"items": items, "errors": errors, "warnings": warnings}


PARSERS = {"equipment": parse_equipment, "workers": parse_workers, "tld": parse_tld,
           "online_doses": parse_online_doses, "inspection": parse_inspection,
           "assessments": parse_assessments}


def _plan_from_legacy_xlsx(data):
    """Your original plan workbook (two header rows, one column per week, ticks)."""
    from . import importers
    from openpyxl import load_workbook
    items, warnings = importers.read_plan(io.BytesIO(data))
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb["برنامه سالانه"]
    title = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))[0] or ""
    m = re.search(r"(1[34]\d\d)", normalize_digits(title))
    out = [{"row": 4 + i, "title": it["title"], "area": it["area"],
            "fields": {"form_title": it["form_title"], "form_code": it["form_code"], "period_weeks": it["period_weeks"]},
            "responsible": it["responsible"], "done_weeks": it["done_weeks"], "plan_year": None}
           for i, it in enumerate(items)]
    return {"items": out, "errors": [], "warnings": [(0, w) for w in warnings],
            "detected_year": int(m.group(1)) if m else None}


def parse_file(kind, filename, data):
    """-> {'items','errors','warnings','detected_year'}; raises ImportFormatError if unreadable."""
    name = (filename or "").lower()
    if kind == "inspection" and name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True)
        if "برنامه سالانه" in wb.sheetnames:
            return _plan_from_legacy_xlsx(data)
    table = load_table(filename, data, kind)
    missing = missing_columns(kind, table["present"])
    if missing:
        raise ImportFormatError("Missing required column(s): " + ", ".join(missing))
    parsed = PARSERS[kind](table)
    parsed["detected_year"] = None
    return parsed


def template_csv(kind):
    headers, examples = TEMPLATES[kind]
    out = io.StringIO()
    w = csv.writer(out)
    w.writerow(headers)
    w.writerows(examples)
    return "\ufeff" + out.getvalue()
