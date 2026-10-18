"""
Pure-Python readers for the legacy Excel files (no Django imports, so they are
easy to test). Each returns plain dicts plus a list of human-readable warnings.
"""
from openpyxl import load_workbook


PERIOD_WEEKS = {"هفتگی": 1, "دوهفته یکبار": 2, "ماهانه": 4, "دو ماهه": 8, "سه ماهه": 12, "شش ماهه": 24}


def clean(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = " ".join(str(v).split())
    return "" if s == "-" else s


# ------------------------------------------------------------------ inspection plan
def read_plan(path, sheet="برنامه سالانه", first_data_row=4, first_week_col=7, weeks=44):
    wb = load_workbook(path, data_only=False)
    ws = wb[sheet]
    warnings, items = [], []
    title = ""
    for n in range(first_data_row, ws.max_row + 1):
        area, period_txt = clean(ws.cell(n, 5).value), clean(ws.cell(n, 6).value)
        if not area and not period_txt:                 # footer / legend starts here
            break
        title = clean(ws.cell(n, 1).value) or title      # titles are merged-style: blank = same as above
        period = PERIOD_WEEKS.get(period_txt)
        if not period:
            warnings.append(f"row {n}: unknown period '{period_txt}' - skipped")
            continue
        done = [w for w in range(1, weeks + 1)
                if clean(ws.cell(n, first_week_col + w - 1).value) == "✓"]
        items.append({
            "title": title, "form_title": clean(ws.cell(n, 2).value),
            "form_code": clean(ws.cell(n, 3).value), "responsible": clean(ws.cell(n, 4).value),
            "area": area, "period_weeks": period, "done_weeks": done,
        })
    # fill a missing form code from another row that uses the same form
    codes = {i["form_title"]: i["form_code"] for i in items if i["form_code"]}
    for i in items:
        if not i["form_code"] and codes.get(i["form_title"]):
            warnings.append(f"'{i['title']} / {i['area']}': form code was empty, filled with {codes[i['form_title']]}")
            i["form_code"] = codes[i["form_title"]]
        elif not i["form_code"]:
            warnings.append(f"'{i['title']} / {i['area']}': form code is empty")
    return items, warnings
