import copy

from docx import Document
from docx.shared import Mm
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from django.utils import timezone


_PERSIAN_DIGITS = str.maketrans("0123456789", "۰۱۲۳۴۵۶۷۸۹")


def _to_persian_digits(text):
    return text.translate(_PERSIAN_DIGITS)


def _gregorian_to_jalali(gy, gm, gd):
    """Convert a Gregorian (y, m, d) to Jalali (Solar Hijri / Shamsi). Pure
    Python, no dependency — same standard algorithm used by most Jalali
    calendar libraries (e.g. jalaali-js), verified against known reference
    dates (Nowruz 2000-03-20 -> 1379/01/01, etc.)."""

    g_d_m = [0, 31, 59, 90, 120, 151, 181, 212, 243, 273, 304, 334]
    if gy > 1600:
        jy = 979
        gy -= 1600
    else:
        jy = 0
        gy -= 621
    gy2 = gy + 1 if gm > 2 else gy
    days = (
        365 * gy
        + (gy2 + 3) // 4
        - (gy2 + 99) // 100
        + (gy2 + 399) // 400
        - 80
        + gd
        + g_d_m[gm - 1]
    )
    jy += 33 * (days // 12053)
    days %= 12053
    jy += 4 * (days // 1461)
    days %= 1461
    if days > 365:
        jy += (days - 1) // 365
        days = (days - 1) % 365
    if days < 186:
        jm = 1 + days // 31
        jd = 1 + (days % 31)
    else:
        jm = 7 + (days - 186) // 30
        jd = 1 + ((days - 186) % 30)
    return jy, jm, jd


def _format_jalali_date(date_obj):
    """Render a Gregorian `date` as a Jalali date string in Persian digits
    (e.g. ۱۴۰۳/۰۲/۱۲), since the form is filled out in Persian."""

    if date_obj is None:
        return ""
    jy, jm, jd = _gregorian_to_jalali(date_obj.year, date_obj.month, date_obj.day)
    return _to_persian_digits(f"{jy:04d}/{jm:02d}/{jd:02d}")


class SpecificationSigner:

    def __init__(self, document_path):

        self.document = Document(document_path)

    def _add_floating_picture(self, paragraph, image_path, width,
                               top_offset=Mm(0), left_offset=Mm(0)):
        """
        Insert `image_path` into `paragraph` as a *floating* picture
        instead of an inline one, so it overlaps the cell on top of
        whatever text is there instead of stretching the row.

        python-docx only inserts pictures inline (part of the text
        flow, which is why Word was growing the row to fit it). We let
        python-docx add the picture normally - it still handles the
        media relationship and sizing - then rewrite the drawing XML it
        produced from an inline <wp:inline> wrapper into a floating
        <wp:anchor> wrapper, which is laid out on its own layer.
        """

        run = paragraph.add_run()
        run.add_picture(image_path, width=width)

        drawing = run._element.find(qn('w:drawing'))
        inline = drawing.find(qn('wp:inline'))

        extent = inline.find(qn('wp:extent'))
        doc_pr = inline.find(qn('wp:docPr'))
        cnv_graphic_frame_pr = inline.find(qn('wp:cNvGraphicFramePr'))
        graphic = inline.find(qn('a:graphic'))

        anchor = OxmlElement('wp:anchor')
        anchor.set('distT', '0')
        anchor.set('distB', '0')
        anchor.set('distL', '0')
        anchor.set('distR', '0')
        anchor.set('simplePos', '0')
        anchor.set('relativeHeight', '251659264')
        anchor.set('behindDoc', '0')  # 0 = drawn in front of the text
        anchor.set('locked', '0')
        anchor.set('layoutInCell', '1')  # keep it anchored inside the table cell
        anchor.set('allowOverlap', '1')

        simple_pos = OxmlElement('wp:simplePos')
        simple_pos.set('x', '0')
        simple_pos.set('y', '0')
        anchor.append(simple_pos)

        position_h = OxmlElement('wp:positionH')
        position_h.set('relativeFrom', 'column')
        h_offset = OxmlElement('wp:posOffset')
        h_offset.text = str(int(left_offset))
        position_h.append(h_offset)
        anchor.append(position_h)

        position_v = OxmlElement('wp:positionV')
        position_v.set('relativeFrom', 'paragraph')
        v_offset = OxmlElement('wp:posOffset')
        v_offset.text = str(int(top_offset))
        position_v.append(v_offset)
        anchor.append(position_v)

        anchor.append(copy.deepcopy(extent))

        effect_extent = OxmlElement('wp:effectExtent')
        for side in ('l', 't', 'r', 'b'):
            effect_extent.set(side, '0')
        anchor.append(effect_extent)

        # wrapNone: text isn't reflowed around the picture, and the
        # picture doesn't add to the space Word reserves for the
        # paragraph/row - this is what keeps the row size unchanged.
        anchor.append(OxmlElement('wp:wrapNone'))
        anchor.append(copy.deepcopy(doc_pr))
        anchor.append(copy.deepcopy(cnv_graphic_frame_pr))
        anchor.append(copy.deepcopy(graphic))

        drawing.replace(inline, anchor)

        return run

    def _get_approval_table(self):
        # The "مدیرعملیات و بهسازی" / "معاون عملیات و بهره‌برداری" grid
        # is a table nested inside the last row of the main table
        # (tables[1], row 16, col 0). Its columns are:
        #   0: row label, 1: نام و نام خانوادگی, 2: تاریخ, 3: امضا
        return self.document.tables[1].cell(16, 0).tables[0]

    def _sign_creator(self, profile):
        if not profile.signature_image:
            raise ValueError("User does not have a signature image.")

        cell = self.document.tables[1].cell(15, 8)

        # paragraphs[0] = "کارشناس هماهنگی و ثبت عملیات:"
        # paragraphs[1] = "تاریخ و امضاء:"  -> date is appended right
        # after this label, on the same line, instead of a new line.
        date_paragraph = cell.paragraphs[1]
        date_paragraph.add_run(
            "  " + _format_jalali_date(timezone.localdate())
        )

        # Floating signature over that same line - no full name and no
        # position are written for this step.
        self._add_floating_picture(
            date_paragraph,
            profile.signature_image.path,
            width=Mm(35),
            top_offset=Mm(-20),
            left_offset=Mm(0),
        )

    def _sign_approval_row(self, row_index, profile):
        if not profile.signature_image:
            raise ValueError("User does not have a signature image.")

        approval_table = self._get_approval_table()

        name_cell = approval_table.cell(row_index, 1)
        date_cell = approval_table.cell(row_index, 2)
        signature_cell = approval_table.cell(row_index, 3)

        name_cell.paragraphs[0].add_run(profile.full_name)
        date_cell.paragraphs[0].add_run(
            _format_jalali_date(timezone.localdate())
        )

        # Floating signature over the "امضا" cell - it does not affect
        # the row's height.
        self._add_floating_picture(
            signature_cell.paragraphs[0],
            profile.signature_image.path,
            width=Mm(28),
            top_offset=Mm(-8),
            left_offset=Mm(0),
        )

    def sign(self, profile, step):

        if step == "CREATOR":
            self._sign_creator(profile)

        elif step == "MANAGER":
            # row 1 of the nested approval table = مدیرعملیات و بهسازی
            self._sign_approval_row(1, profile)

        elif step == "DEPUTY":
            # row 2 of the nested approval table = معاون عملیات و بهره‌برداری
            self._sign_approval_row(2, profile)
        
        elif step == "CEO":
            self._sign_approval_row(3, profile)

        else:
            raise ValueError("Unknown approval step.")

    def save(self, filename):

        self.document.save(filename)
