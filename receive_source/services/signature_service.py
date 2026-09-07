"""
Stamps signatures onto the generated receive-specification .docx.

This is a SEPARATE class from operations.services.signature_service.
SpecificationSigner — that one is hardcoded to the LICENSE template's
row 15/16 positions and only knows 3 steps (CREATOR/MANAGER/DEPUTY).
The receive template has a different layout (verified against the
attached blank template's raw XML):

  row 21: creator signature — flat row (not nested):
          col0-2 label "کارشناس هماهنگی و ثبت عملیات"
          col3-10 name, col11-16 date, col17-20 signature

  row 22, cell(0): contains a NESTED table with rows for the remaining
          four signers:
          row 0: header (blank / نام و نام خانوادگی / تاریخ / امضا)
          row 1: مدیرعملیات و بهسازی           -> MANAGER
          row 2: مدیرکنترل عملیات پرتوی و صنعتی -> CONTROL
          row 3: معاون عملیات و بهره‌برداری      -> DEPUTY
          row 4: مدیرعامل (در صورت اعمال تخفیف) -> CEO
          cols: 0 role label, 1 name, 2 date, 3 signature

The floating-picture insertion logic is copied as-is from the license
app's signer (it's generic OOXML manipulation, not template-specific).
"""

import copy

from docx.shared import Mm
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from django.utils import timezone

from .specification_generator import _format_jalali_date


class ReceiveSpecificationSigner:

    def __init__(self, document_path):
        from docx import Document
        self.document = Document(document_path)

    def _add_floating_picture(self, paragraph, image_path, width,
                               top_offset=Mm(0), left_offset=Mm(0)):
        """Ported verbatim from operations.services.signature_service —
        generic OOXML manipulation to turn an inline picture into a
        floating one that overlaps existing cell text instead of
        stretching the row."""

        run = paragraph.add_run()
        run.add_picture(image_path, width=width)

        drawing = run._element.find(qn('w:drawing'))
        inline = drawing.find(qn('wp:inline'))

        extent = inline.find(qn('wp:extent'))
        doc_pr = inline.find(qn('wp:docPr'))
        cnv_graphic_frame_pr = inline.find(qn('wp:cNvGraphicFramePr'))
        graphic = inline.find(qn('a:graphic'))

        anchor = OxmlElement('wp:anchor')
        anchor.set('distT', '0')
        anchor.set('distB', '0')
        anchor.set('distL', '0')
        anchor.set('distR', '0')
        anchor.set('simplePos', '0')
        anchor.set('relativeHeight', '251659264')
        anchor.set('behindDoc', '0')
        anchor.set('locked', '0')
        anchor.set('layoutInCell', '1')
        anchor.set('allowOverlap', '1')

        simple_pos = OxmlElement('wp:simplePos')
        simple_pos.set('x', '0')
        simple_pos.set('y', '0')
        anchor.append(simple_pos)

        position_h = OxmlElement('wp:positionH')
        position_h.set('relativeFrom', 'column')
        h_offset = OxmlElement('wp:posOffset')
        h_offset.text = str(int(left_offset))
        position_h.append(h_offset)
        anchor.append(position_h)

        position_v = OxmlElement('wp:positionV')
        position_v.set('relativeFrom', 'paragraph')
        v_offset = OxmlElement('wp:posOffset')
        v_offset.text = str(int(top_offset))
        position_v.append(v_offset)
        anchor.append(position_v)

        anchor.append(copy.deepcopy(extent))

        effect_extent = OxmlElement('wp:effectExtent')
        for side in ('l', 't', 'r', 'b'):
            effect_extent.set(side, '0')
        anchor.append(effect_extent)

        anchor.append(OxmlElement('wp:wrapNone'))
        anchor.append(copy.deepcopy(doc_pr))
        anchor.append(copy.deepcopy(cnv_graphic_frame_pr))
        anchor.append(copy.deepcopy(graphic))

        drawing.replace(inline, anchor)

        return run

    def _get_approval_table(self):
        return self.document.tables[1].cell(22, 0).tables[0]

    def _sign_creator(self, profile):
        if not profile.signature_image:
            raise ValueError("User does not have a signature image.")

        table = self.document.tables[1]

        name_cell = table.cell(21, 3)
        date_cell = table.cell(21, 11)
        signature_cell = table.cell(21, 17)

        name_cell.paragraphs[0].add_run(profile.full_name)
        date_cell.paragraphs[0].add_run(_format_jalali_date(timezone.localdate()))

        self._add_floating_picture(
            signature_cell.paragraphs[0],
            profile.signature_image.path,
            width=Mm(28),
            top_offset=Mm(-8),
            left_offset=Mm(0),
        )

    def _sign_approval_row(self, row_index, profile):
        if not profile.signature_image:
            raise ValueError("User does not have a signature image.")

        approval_table = self._get_approval_table()

        name_cell = approval_table.cell(row_index, 1)
        date_cell = approval_table.cell(row_index, 2)
        signature_cell = approval_table.cell(row_index, 3)

        name_cell.paragraphs[0].add_run(profile.full_name)
        date_cell.paragraphs[0].add_run(_format_jalali_date(timezone.localdate()))

        self._add_floating_picture(
            signature_cell.paragraphs[0],
            profile.signature_image.path,
            width=Mm(28),
            top_offset=Mm(-8),
            left_offset=Mm(0),
        )

    # Row indices inside the nested approval table (row22, cell0, tables[0]).
    _APPROVAL_ROW_BY_STEP = {
        "MANAGER": 1,
        "CONTROL": 2,
        "DEPUTY": 3,
        "CEO": 4,
    }

    def sign(self, profile, step):

        if step == "CREATOR":
            self._sign_creator(profile)
            return

        row_index = self._APPROVAL_ROW_BY_STEP.get(step)
        if row_index is None:
            raise ValueError(f"Unknown approval step: {step}")

        self._sign_approval_row(row_index, profile)

    def save(self, filename):
        self.document.save(filename)
