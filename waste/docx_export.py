"""
Fill the two official minutes forms (C2-FRM-03 solid / C2-FRM-04 liquid) from a
plain-dict context. Deliberately Django-free so it can be unit-tested; the
Django side only has to build the context (see services.receipt_to_context).

Layout facts this relies on (checked against the two uploaded templates):
  SOLID  table1 header paragraphs | table2 rows 2-8 = material rows, row 9 = total,
         cols 1 vol, 2 mass, 3 bag, 4 bin, 5 box, 6 drum, 7 container, 8 non-standard,
         9 surface dose, 10 half-life | table3 detail cells | table4 signatures
  LIQUID table1 header paragraphs | table2 rows 2-3 = half-life classes,
         rows 4-6 = activity bands; cols 1 vol, 2 alpha, 3 beta, 4 gamma, 5 pH,
         6 density, 7 hardness | table3 one cell, label paragraphs | table4 signatures
"""
import copy
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

TEMPLATE_DIR = Path(__file__).parent / "docx_templates"
SOLID_TEMPLATE = TEMPLATE_DIR / "solid_minutes_C2-FRM-03.docx"
LIQUID_TEMPLATE = TEMPLATE_DIR / "liquid_minutes_C2-FRM-04.docx"

# form row index for each MaterialType value (solid form)
SOLID_ROW = {
    "LIGHTWEIGHT": 2, "HEAVY": 3, "BIOLOGICAL": 4, "FILTERS": 5, "RESINS": 6,
    "SOIL_SEDIMENTS_SLUDGE": 7,
    "OTHER": 8, "SPECIAL_OTHER": 8, "CHARCOAL": 8,       # the form has no own row for these
}
TOTAL_ROW = 9
# form column index for each PackageType value (solid form)
SOLID_PKG_COL = {
    "BAG": 3, "SMALL_BIN": 4, "LARGE_BIN": 4, "BOX": 5, "DRUM": 6, "CONTAINER": 7,
    "NON_STANDARD": 8, "CAN": 8, "SMALL_PIPE": 8,          # no own column on the form
}
LIQ_HALFLIFE_ROW = {"LT100": 2, "GT100": 3}
LIQ_ACTIVITY_ROW = {"LOW": 4, "MID": 5, "HIGH": 6}        # <8000, 8000-100000, >100000 Bq/L


# ----------------------------------------------------------------- helpers
def _fmt(v, nd=3):
    if v is None or v == "":
        return ""
    try:
        f = float(v)
    except (TypeError, ValueError):
        return str(v)
    return f"{f:.{nd}f}".rstrip("0").rstrip(".") if f % 1 else str(int(f))


def _set_par(p, text):
    """Replace a paragraph's text but keep the formatting of its first run
    (or of the paragraph mark when the paragraph has no runs yet)."""
    if p.runs:
        p.runs[0].text = text
        for r in p.runs[1:]:
            r._r.getparent().remove(r._r)
        return
    run = p.add_run(text)
    mark = p._p.find(qn("w:pPr") + "/" + qn("w:rPr"))
    if mark is not None:
        run._r.insert(0, copy.deepcopy(mark))


def _set_cell(cell, text):
    _set_par(cell.paragraphs[0], str(text))


def _label_par(cell, label, value):
    """Find the paragraph that starts with `label` and write 'label value'."""
    for p in cell.paragraphs:
        if p.text.strip().startswith(label):
            _set_par(p, f"{label} {value}")
            return True
    return False


def _fill_header(doc, ctx):
    t0 = doc.tables[0]
    c = t0.cell(0, 0)
    _label_par(c, "تاریخ :", ctx.get("minutes_date", ""))
    _label_par(c, "شماره :", ctx.get("minutes_number", ""))

    cell = doc.tables[1].cell(0, 0)
    gap = "          "
    for p in cell.paragraphs:
        t = p.text.strip()
        if t.startswith("مرکز تحویل دهنده"):
            _set_par(p, f"مرکز تحویل دهنده: {ctx.get('origin_facility', '')}{gap}آزمایشگاه: {ctx.get('laboratory', '')}")
        elif t.startswith("تاریخ درخواست"):
            _set_par(p, f"تاریخ درخواست تحویل پسماند: {ctx.get('request_date', '')}{gap}شماره درخواست تحویل پسماند: {ctx.get('request_number', '')}")
        elif t.startswith("تاریخ تحویل"):
            _set_par(p, f"تاریخ تحویل پسماند: {ctx.get('delivery_date', '')}")


def _fill_signatures(doc, ctx):
    t = doc.tables[4]
    for row_i, key in ((1, "deliverer"), (2, "receiver")):
        who = ctx.get(key, {})
        _set_cell(t.cell(row_i, 1), who.get("name", ""))
        _set_cell(t.cell(row_i, 2), who.get("position", ""))
        _set_cell(t.cell(row_i, 3), who.get("date", ""))
        # column 4 (signature) is intentionally left blank: it is signed on paper


def _fill_details_solid(doc, ctx):
    t = doc.tables[3]
    _set_cell(t.cell(0, 1), ctx.get("radionuclides", ""))
    _set_cell(t.cell(1, 1), ctx.get("waste_origin_place", ""))
    _set_cell(t.cell(2, 1), ctx.get("cabin_dose", ""))
    _set_cell(t.cell(2, 3), ctx.get("container_dose", ""))
    _set_cell(t.cell(3, 1), ctx.get("description", ""))


def _fill_details_liquid(doc, ctx):
    cell = doc.tables[3].cell(0, 0)
    _label_par(cell, "رادیونوکلیدهای موجود در پسماند:", ctx.get("radionuclides", ""))
    _label_par(cell, "محل تولید پسماند:", ctx.get("waste_origin_place", ""))
    for p in cell.paragraphs:
        if p.text.strip().startswith("دز کابین"):
            _set_par(p, f"دز کابین راننده: {ctx.get('cabin_dose', '')}          "
                        f"دز کانتینر حامل پسماند: {ctx.get('container_dose', '')}")
    _label_par(cell, "توضیحات:", ctx.get("description", ""))


# ------------------------------------------------------------------ solid
def fill_solid(ctx, out_path):
    """ctx['lines']: dicts with material, package_type, package_count, mass_kg,
    volume_m3, surface_dose, half_life. Several lines may share one material row
    (different packages) - volume/mass are summed, package counts go in their column."""
    doc = Document(str(SOLID_TEMPLATE))
    _fill_header(doc, ctx)
    t = doc.tables[2]

    rows, total = {}, {"vol": 0.0, "mass": 0.0, "pk": {}, "dose": None}
    for ln in ctx["lines"]:
        r = SOLID_ROW[ln["material"]]
        d = rows.setdefault(r, {"vol": 0.0, "mass": 0.0, "pk": {}, "dose": None, "hl": []})
        vol, mass = float(ln.get("volume_m3") or 0), float(ln.get("mass_kg") or 0)
        d["vol"] += vol; d["mass"] += mass
        total["vol"] += vol; total["mass"] += mass
        col = SOLID_PKG_COL[ln["package_type"]]
        n = int(ln.get("package_count") or 1)
        d["pk"][col] = d["pk"].get(col, 0) + n
        total["pk"][col] = total["pk"].get(col, 0) + n
        dose = ln.get("surface_dose")
        if dose is not None:
            d["dose"] = max(float(dose), d["dose"] or 0)
            total["dose"] = max(float(dose), total["dose"] or 0)
        if ln.get("half_life") and ln["half_life"] not in d["hl"]:
            d["hl"].append(ln["half_life"])

    def write(r, d, hl=""):
        _set_cell(t.cell(r, 1), _fmt(d["vol"]) if d["vol"] else "")
        _set_cell(t.cell(r, 2), _fmt(d["mass"]) if d["mass"] else "")
        for col, n in d["pk"].items():
            _set_cell(t.cell(r, col), n)
        _set_cell(t.cell(r, 9), _fmt(d["dose"]) if d["dose"] is not None else "")
        _set_cell(t.cell(r, 10), hl)

    for r, d in rows.items():
        write(r, d, " / ".join(d["hl"]))
    write(TOTAL_ROW, total)

    _fill_details_solid(doc, ctx)
    _fill_signatures(doc, ctx)
    doc.save(out_path)
    return out_path


# ----------------------------------------------------------------- liquid
def liquid_activity_class(total_bq_per_l):
    if total_bq_per_l is None:
        return None
    return "LOW" if total_bq_per_l < 8000 else ("MID" if total_bq_per_l <= 100000 else "HIGH")


def fill_liquid(ctx, out_path):
    """ctx['lines']: dicts with half_life_class (LT100/GT100), volume_m3, alpha, beta,
    gamma (Bq/L), ph, density, hardness. The form is a classification matrix, so each
    stream is written into BOTH its half-life row and its activity-band row
    (band is derived from alpha+beta+gamma)."""
    doc = Document(str(LIQUID_TEMPLATE))
    _fill_header(doc, ctx)
    t = doc.tables[2]

    def put(r, ln):
        # if a row already holds a stream, append ' / ' so nothing is overwritten
        for col, key in ((1, "volume_m3"), (2, "alpha"), (3, "beta"), (4, "gamma"),
                         (5, "ph"), (6, "density"), (7, "hardness")):
            new = _fmt(ln.get(key))
            old = t.cell(r, col).text.strip()
            _set_cell(t.cell(r, col), f"{old} / {new}" if old and new else (old or new))

    for ln in ctx["lines"]:
        if ln.get("half_life_class") in LIQ_HALFLIFE_ROW:
            put(LIQ_HALFLIFE_ROW[ln["half_life_class"]], ln)
        parts = [ln.get(k) for k in ("alpha", "beta", "gamma")]
        known = [float(x) for x in parts if x is not None]
        band = liquid_activity_class(sum(known)) if known else None
        if band:
            put(LIQ_ACTIVITY_ROW[band], ln)

    _fill_details_liquid(doc, ctx)
    _fill_signatures(doc, ctx)
    doc.save(out_path)
    return out_path